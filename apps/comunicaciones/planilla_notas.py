"""Notas de toda la gestión en Excel: una hoja por curso.

Cada hoja se lee sola, sin ir al sistema: arriba el curso y sus cifras, en el
medio las notas de cada estudiante por materia y trimestre, y abajo cada
materia con su docente y cuántas notas lleva cargadas, que es lo que dice si un
promedio bajo es rendimiento o simplemente notas que faltan.

La primera hoja compara los cursos entre sí.

Los promedios salen de `pivotar_notas`, la misma cuenta que ve el estudiante en
«Mis notas»: si el Excel y la pantalla dieran números distintos, nadie sabría
cuál creer.
"""
import collections

from django.utils import formats, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.academico.models import AsignacionDocente, Curso, Matricula
from apps.asistencia.models import Asistencia
from apps.asistencia.planillas import nombre_de_hoja
from apps.calificaciones.models import NOTA_APROBACION, Nota
from apps.calificaciones.services import pivotar_notas

from .reportes import COLEGIO, NIVEL

FILA_MATERIAS = 8
FILA_SUBTITULOS = 9
PRIMERA_FILA = 10

FIJAS = ['N°', 'RUDE', 'Estudiante']
POR_MATERIA = [('t1', '1.er trim.'), ('t2', '2.º trim.'), ('t3', '3.er trim.'), ('promedio', 'Promedio')]
DEL_ESTUDIANTE = [('t1', 'Prom. 1.er trim.'), ('t2', 'Prom. 2.º trim.'),
                  ('t3', 'Prom. 3.er trim.'), ('promedio', 'Nota total')]
FINALES = ['Materias reprobadas', 'Estado', 'Faltas', 'Atrasos']

APROBANDO, REPROBANDO, SIN_NOTAS = 'Aprobando', 'Reprobando', 'Sin notas'

CENTRADO = Alignment(horizontal='center', vertical='center', wrap_text=True)
IZQUIERDA = Alignment(horizontal='left', vertical='center')
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
APAGADO = Font(italic=True, color='666666')
ROJO_TEXTO = Font(color='9C2B23', bold=True)
FONDO_TITULOS = PatternFill('solid', fgColor='EDF6F4')
FONDO_MATERIA = PatternFill('solid', fgColor='D8EEEA')
FONDO_RESUMEN = PatternFill('solid', fgColor='FFF4DC')
FONDO_REPROBADO = PatternFill('solid', fgColor='F8D8D4')
FONDO_APROBADO = PatternFill('solid', fgColor='D6F0E4')


def _numero(valor):
    """`pivotar_notas` marca con '-' lo que no hay; en la planilla va vacío."""
    return None if valor in (None, '-') else round(float(valor), 2)


def _promedio(valores):
    valores = [v for v in valores if v is not None]
    return round(sum(valores) / len(valores), 2) if valores else None


def _nota(celda, valor):
    """Escribe una nota; las que no alcanzan para aprobar quedan en rojo."""
    celda.alignment = CENTRADO
    if valor is None:
        celda.value = '—'
        celda.font = APAGADO
        return
    celda.value = valor
    if valor < NOTA_APROBACION:
        celda.fill = FONDO_REPROBADO
        celda.font = ROJO_TEXTO


def _titulo(hoja, fila, columna, texto, fondo=FONDO_TITULOS, hasta_fila=None, hasta_columna=None):
    celda = hoja.cell(row=fila, column=columna, value=texto)
    celda.font = NEGRITA
    celda.alignment = CENTRADO
    celda.fill = fondo
    if hasta_fila or hasta_columna:
        hoja.merge_cells(
            start_row=fila, start_column=columna,
            end_row=hasta_fila or fila, end_column=hasta_columna or columna,
        )
    return celda


def _generado(gestion):
    hoy = formats.date_format(timezone.localdate(), 'j \\d\\e F \\d\\e Y')
    return f'Gestión {gestion.anio} · generado el {hoy}'


def _cifras(hoja, fila, pares):
    """Una fila de rótulos y debajo sus valores: lo primero que se mira."""
    for columna, (rotulo, valor) in enumerate(pares, start=1):
        _titulo(hoja, fila, columna, rotulo)
        celda = hoja.cell(row=fila + 1, column=columna, value=valor if valor is not None else '—')
        celda.alignment = CENTRADO
        celda.font = Font(bold=True, size=12)


# ─────────────────────────────── datos ───────────────────────────────

def _cargar(gestion):
    """Todo en pocas consultas: son unos 400 estudiantes y miles de notas."""
    matriculas = collections.defaultdict(list)
    for m in (
        Matricula.objects.filter(activa=True, curso__gestion=gestion)
        .select_related('estudiante__usuario')
        .order_by('estudiante__usuario__last_name', 'estudiante__usuario__first_name')
    ):
        matriculas[m.curso_id].append(m)

    asignaciones = collections.defaultdict(list)
    for a in AsignacionDocente.objects.filter(curso__gestion=gestion).select_related('materia', 'profesor'):
        asignaciones[a.curso_id].append(a)

    notas = collections.defaultdict(list)
    for n in Nota.objects.calificadas().filter(asignacion__curso__gestion=gestion).select_related('asignacion__materia'):
        notas[n.estudiante_id].append(n)

    asistencia = collections.defaultdict(collections.Counter)
    for estudiante_id, estado in (
        Asistencia.objects.filter(fecha__year=gestion.anio).values_list('estudiante_id', 'estado')
    ):
        asistencia[estudiante_id][estado] += 1

    return matriculas, asignaciones, notas, asistencia


def _filas_del_curso(matriculas, notas, asistencia):
    filas = []
    for matricula in matriculas:
        estudiante = matricula.estudiante
        suyas = notas.get(estudiante.pk, [])
        por_materia_filas, totales = pivotar_notas(suyas)
        por_materia = {
            f['materia']: {clave: _numero(f[clave]) for clave, _ in POR_MATERIA}
            for f in por_materia_filas
        }
        reprobadas = [m for m, v in por_materia.items() if v['promedio'] is not None and v['promedio'] < NOTA_APROBACION]
        if not suyas:
            estado = SIN_NOTAS
        elif reprobadas:
            estado = REPROBANDO
        else:
            estado = APROBANDO
        filas.append({
            'estudiante': estudiante,
            'por_materia': por_materia,
            'notas': suyas,
            'totales': {clave: _numero((totales or {}).get(clave)) for clave, _ in DEL_ESTUDIANTE},
            'reprobadas': reprobadas,
            'estado': estado,
            'faltas': asistencia[estudiante.pk][Asistencia.Estado.FALTA],
            'atrasos': asistencia[estudiante.pk][Asistencia.Estado.ATRASO],
        })
    return filas


# ─────────────────────────────── hojas ───────────────────────────────

def _hoja_del_curso(libro, nombre_hoja, curso, gestion, matriculas, asignaciones, notas, asistencia):
    hoja = libro.create_sheet(nombre_hoja)
    filas = _filas_del_curso(matriculas, notas, asistencia)

    # Las materias del curso aunque todavía no tengan notas, y alguna nota
    # suelta de una materia que ya no esté asignada, para no perderla.
    docentes = {a.materia.nombre: a.profesor.get_full_name() for a in asignaciones}
    materias = sorted(set(docentes) | {m for f in filas for m in f['por_materia']})

    # ── cabecera ──
    hoja['A1'] = f'{COLEGIO} · {NIVEL}'
    hoja['A1'].font = APAGADO
    hoja['A2'] = f'Notas de {curso}'
    hoja['A2'].font = TITULO
    hoja['A3'] = _generado(gestion)
    hoja['A3'].font = APAGADO

    totales = [f['totales']['promedio'] for f in filas]
    estados = collections.Counter(f['estado'] for f in filas)
    con_nota = [(f['totales']['promedio'], f['estudiante']) for f in filas if f['totales']['promedio'] is not None]
    mejor = max(con_nota, key=lambda x: x[0]) if con_nota else None
    _cifras(hoja, 5, [
        ('Estudiantes', len(filas)),
        ('Promedio del curso', _promedio(totales)),
        (APROBANDO, estados[APROBANDO]),
        (REPROBANDO, estados[REPROBANDO]),
        (SIN_NOTAS, estados[SIN_NOTAS]),
        ('Faltas del curso', sum(f['faltas'] for f in filas)),
        ('Mejor promedio', f'{mejor[1].usuario.get_full_name()} ({mejor[0]})' if mejor else None),
    ])

    # ── títulos de la tabla ──
    for columna, texto in enumerate(FIJAS, start=1):
        _titulo(hoja, FILA_MATERIAS, columna, texto, hasta_fila=FILA_SUBTITULOS)

    columna = len(FIJAS) + 1
    inicio_materia = {}
    for materia in materias:
        inicio_materia[materia] = columna
        encabezado = materia if materia not in docentes else f'{materia}\n{docentes[materia]}'
        _titulo(hoja, FILA_MATERIAS, columna, encabezado, FONDO_MATERIA, hasta_columna=columna + len(POR_MATERIA) - 1)
        for desplazamiento, (_clave, texto) in enumerate(POR_MATERIA):
            _titulo(hoja, FILA_SUBTITULOS, columna + desplazamiento, texto, FONDO_MATERIA)
        columna += len(POR_MATERIA)

    inicio_resumen = columna
    _titulo(hoja, FILA_MATERIAS, columna, 'Resumen del estudiante', FONDO_RESUMEN,
            hasta_columna=columna + len(DEL_ESTUDIANTE) + len(FINALES) - 1)
    for desplazamiento, texto in enumerate([t for _c, t in DEL_ESTUDIANTE] + FINALES):
        _titulo(hoja, FILA_SUBTITULOS, columna + desplazamiento, texto, FONDO_RESUMEN)
    ultima_columna = columna + len(DEL_ESTUDIANTE) + len(FINALES) - 1
    hoja.row_dimensions[FILA_MATERIAS].height = 48

    # ── estudiantes ──
    for indice, fila in enumerate(filas):
        r = PRIMERA_FILA + indice
        hoja.cell(row=r, column=1, value=indice + 1).alignment = CENTRADO
        hoja.cell(row=r, column=2, value=fila['estudiante'].rude).alignment = CENTRADO
        hoja.cell(row=r, column=3, value=fila['estudiante'].usuario.get_full_name()).alignment = IZQUIERDA

        for materia in materias:
            valores = fila['por_materia'].get(materia, {})
            for desplazamiento, (clave, _texto) in enumerate(POR_MATERIA):
                celda = hoja.cell(row=r, column=inicio_materia[materia] + desplazamiento)
                _nota(celda, valores.get(clave))
                if clave == 'promedio':
                    celda.font = Font(bold=True, italic=celda.font.italic, color=celda.font.color)

        c = inicio_resumen
        for desplazamiento, (clave, _texto) in enumerate(DEL_ESTUDIANTE):
            _nota(hoja.cell(row=r, column=c + desplazamiento), fila['totales'][clave])
        c += len(DEL_ESTUDIANTE)

        reprobadas = hoja.cell(row=r, column=c, value=len(fila['reprobadas']))
        reprobadas.alignment = CENTRADO
        if fila['reprobadas']:
            reprobadas.value = f"{len(fila['reprobadas'])}: {', '.join(fila['reprobadas'])}"
            reprobadas.alignment = Alignment(vertical='center', wrap_text=True)
        estado = hoja.cell(row=r, column=c + 1, value=fila['estado'])
        estado.alignment = CENTRADO
        if fila['estado'] == REPROBANDO:
            estado.fill, estado.font = FONDO_REPROBADO, ROJO_TEXTO
        elif fila['estado'] == APROBANDO:
            estado.fill = FONDO_APROBADO
        else:
            estado.font = APAGADO
        hoja.cell(row=r, column=c + 2, value=fila['faltas']).alignment = CENTRADO
        hoja.cell(row=r, column=c + 3, value=fila['atrasos']).alignment = CENTRADO

    # ── pie: cómo va el curso en cada columna ──
    r = PRIMERA_FILA + len(filas)
    if filas:
        hoja.cell(row=r, column=3, value='Promedio del curso').font = NEGRITA
        hoja.cell(row=r + 1, column=3, value='Reprobados').font = NEGRITA
        for materia in materias:
            for desplazamiento, (clave, _texto) in enumerate(POR_MATERIA):
                valores = [f['por_materia'].get(materia, {}).get(clave) for f in filas]
                _nota(hoja.cell(row=r, column=inicio_materia[materia] + desplazamiento), _promedio(valores))
                abajo = [v for v in valores if v is not None and v < NOTA_APROBACION]
                celda = hoja.cell(row=r + 1, column=inicio_materia[materia] + desplazamiento, value=len(abajo))
                celda.alignment = CENTRADO
        for desplazamiento, (clave, _texto) in enumerate(DEL_ESTUDIANTE):
            valores = [f['totales'][clave] for f in filas]
            _nota(hoja.cell(row=r, column=inicio_resumen + desplazamiento), _promedio(valores))
            abajo = [v for v in valores if v is not None and v < NOTA_APROBACION]
            hoja.cell(row=r + 1, column=inicio_resumen + desplazamiento, value=len(abajo)).alignment = CENTRADO
        r += 3
    else:
        hoja.cell(row=r, column=3, value='Este curso no tiene estudiantes matriculados.').font = APAGADO
        r += 2

    # ── materias: docente y notas cargadas ──
    hoja.cell(row=r, column=3, value='Materias del curso').font = Font(bold=True, size=12)
    r += 1
    titulos = ['Materia', 'Docente', 'Promedio', 'Reprobados',
               'Prom. ser', 'Prom. saber', 'Prom. hacer', 'Prom. autoev.', 'Prom. extra',
               'Notas 1.er trim.', 'Notas 2.º trim.', 'Notas 3.er trim.']
    for desplazamiento, texto in enumerate(titulos):
        _titulo(hoja, r, 3 + desplazamiento, texto)
    resumen_materias = []
    for materia in materias:
        r += 1
        promedios = [f['por_materia'].get(materia, {}).get('promedio') for f in filas]
        promedio = _promedio(promedios)
        abajo = sum(1 for v in promedios if v is not None and v < NOTA_APROBACION)
        hoja.cell(row=r, column=3, value=materia)
        hoja.cell(row=r, column=4, value=docentes.get(materia, 'Sin docente asignado'))
        _nota(hoja.cell(row=r, column=5), promedio)
        hoja.cell(row=r, column=6, value=abajo).alignment = CENTRADO

        # Promedio del curso en cada parte de la nota: dice dónde está el problema.
        de_la_materia = [n for f in filas for n in f['notas'] if n.asignacion.materia.nombre == materia]
        for desplazamiento, parte in enumerate(['ser', 'saber', 'hacer', 'autoevaluacion', 'extracurricular']):
            valores = [float(getattr(n, parte)) for n in de_la_materia if getattr(n, parte) is not None]
            celda = hoja.cell(row=r, column=7 + desplazamiento,
                              value=round(sum(valores) / len(valores), 2) if valores else '—')
            celda.alignment = CENTRADO

        cargadas = []
        for desplazamiento, clave in enumerate(['t1', 't2', 't3']):
            cuantas = sum(1 for f in filas if f['por_materia'].get(materia, {}).get(clave) is not None)
            cargadas.append(cuantas)
            celda = hoja.cell(row=r, column=12 + desplazamiento, value=f'{cuantas} de {len(filas)}')
            celda.alignment = CENTRADO
            if filas and 0 < cuantas < len(filas):
                celda.fill = FONDO_RESUMEN      # a medio cargar: alguien tiene que completarla
        resumen_materias.append((materia, promedio, cargadas))

    # ── anchos ──
    hoja.column_dimensions['A'].width = 5
    hoja.column_dimensions['B'].width = 18
    hoja.column_dimensions['C'].width = 34
    for columna in range(len(FIJAS) + 1, ultima_columna + 1):
        hoja.column_dimensions[get_column_letter(columna)].width = 10
    for desplazamiento in range(len(DEL_ESTUDIANTE)):
        hoja.column_dimensions[get_column_letter(inicio_resumen + desplazamiento)].width = 12
    hoja.column_dimensions[get_column_letter(inicio_resumen + len(DEL_ESTUDIANTE))].width = 28
    hoja.column_dimensions[get_column_letter(inicio_resumen + len(DEL_ESTUDIANTE) + 1)].width = 13
    # La tabla de materias usa la columna D para el docente: que se lea el nombre.
    hoja.column_dimensions['D'].width = max(hoja.column_dimensions['D'].width or 10, 30)
    # Número, RUDE y nombre quedan a la vista al desplazarse por las materias.
    hoja.freeze_panes = hoja.cell(row=PRIMERA_FILA, column=len(FIJAS) + 1)

    return {
        'curso': curso,
        'hoja': nombre_hoja,
        'estudiantes': len(filas),
        'promedio': _promedio(totales),
        'estados': estados,
        'materias': resumen_materias,
        'faltas': sum(f['faltas'] for f in filas),
        'atrasos': sum(f['atrasos'] for f in filas),
    }


def _hoja_resumen(hoja, gestion, cursos):
    hoja.title = 'Resumen'
    hoja['A1'] = f'{COLEGIO} · {NIVEL}'
    hoja['A1'].font = APAGADO
    hoja['A2'] = 'Notas por curso'
    hoja['A2'].font = TITULO
    hoja['A3'] = _generado(gestion)
    hoja['A3'].font = APAGADO

    todos = sum(c['estudiantes'] for c in cursos)
    estados = sum((c['estados'] for c in cursos), collections.Counter())
    _cifras(hoja, 5, [
        ('Estudiantes', todos),
        ('Promedio del colegio', _promedio([c['promedio'] for c in cursos])),
        (APROBANDO, estados[APROBANDO]),
        (REPROBANDO, estados[REPROBANDO]),
        (SIN_NOTAS, estados[SIN_NOTAS]),
    ])

    fila = 8
    titulos = ['Curso', 'Estudiantes', 'Promedio', APROBANDO, REPROBANDO, SIN_NOTAS,
               'Materia más baja', 'Notas 1.er trim.', 'Notas 2.º trim.', 'Notas 3.er trim.', 'Faltas', 'Atrasos']
    for columna, texto in enumerate(titulos, start=1):
        _titulo(hoja, fila, columna, texto)

    for datos in cursos:
        fila += 1
        hoja.cell(row=fila, column=1, value=datos['hoja']).font = NEGRITA
        hoja.cell(row=fila, column=2, value=datos['estudiantes']).alignment = CENTRADO
        _nota(hoja.cell(row=fila, column=3), datos['promedio'])
        for columna, estado in enumerate([APROBANDO, REPROBANDO, SIN_NOTAS], start=4):
            hoja.cell(row=fila, column=columna, value=datos['estados'][estado]).alignment = CENTRADO

        con_promedio = [(p, m) for m, p, _c in datos['materias'] if p is not None]
        baja = min(con_promedio) if con_promedio else None
        hoja.cell(row=fila, column=7, value=f'{baja[1]} ({baja[0]})' if baja else '—')

        # Porcentaje de notas cargadas: estudiantes × materias posibles.
        posibles = datos['estudiantes'] * len(datos['materias'])
        for desplazamiento in range(3):
            cargadas = sum(c[desplazamiento] for _m, _p, c in datos['materias'])
            celda = hoja.cell(row=fila, column=8 + desplazamiento, value=(cargadas / posibles) if posibles else None)
            celda.number_format = '0%'
            celda.alignment = CENTRADO
        hoja.cell(row=fila, column=11, value=datos['faltas']).alignment = CENTRADO
        hoja.cell(row=fila, column=12, value=datos['atrasos']).alignment = CENTRADO

    anchos = [14, 12, 11, 12, 12, 11, 36, 14, 14, 14, 9, 9]
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    hoja.freeze_panes = 'B9'


def libro_de_notas(gestion):
    """Libro completo: el resumen de todos los cursos y después una hoja por curso."""
    libro = Workbook()
    resumen = libro.active
    matriculas, asignaciones, notas, asistencia = _cargar(gestion)

    cursos = list(Curso.objects.filter(gestion=gestion))
    niveles = {c.nivel for c in cursos}
    datos = []
    for curso in cursos:
        corto = f'{curso.grado}° {curso.paralelo}'
        if len(niveles) > 1:
            corto = f'{curso.get_nivel_display()[:3]} {corto}'
        datos.append(_hoja_del_curso(
            libro, nombre_de_hoja(corto), curso, gestion,
            matriculas.get(curso.pk, []), asignaciones.get(curso.pk, []), notas, asistencia,
        ))

    _hoja_resumen(resumen, gestion, datos)
    return libro


# ══════════════ la planilla del docente, tal como la llena ══════════════

def libro_de_planilla(asignacion, trimestre, matriculas, actividades, puntajes, notas):
    """Un Excel con la planilla de una materia y un trimestre.

    Es la hoja de papel del docente: sus actividades en columnas, agrupadas por
    dimensión, con el promedio de cada una, la autoevaluación, lo extracurricular
    y la nota final.
    """
    from apps.calificaciones.models import (
        MAXIMO, MAXIMO_AUTOEVALUACION, MAXIMO_EXTRACURRICULAR, Dimension, Trimestre,
    )

    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_de_hoja(f'{asignacion.curso.grado}-{asignacion.curso.paralelo} T{trimestre}')

    hoja['A1'] = f'{COLEGIO} · {NIVEL}'
    hoja['A1'].font = APAGADO
    hoja['A2'] = f'{asignacion.materia} · {asignacion.curso}'
    hoja['A2'].font = TITULO
    hoja['A3'] = (f'{Trimestre(trimestre).label} · {asignacion.profesor.get_full_name()} · '
                  f'{_generado(asignacion.curso.gestion)}')
    hoja['A3'].font = APAGADO
    hoja['A4'] = ('Ser sobre 10 · saber sobre 45 · hacer sobre 40 · autoevaluación sobre 5. '
                  'Lo extracurricular suma hasta 5 puntos, sin pasar de 100.')
    hoja['A4'].font = APAGADO

    fila_grupos, fila_titulos, primera = 6, 7, 8
    for columna, fijo in enumerate(['N°', 'RUDE', 'Estudiante'], start=1):
        _titulo(hoja, fila_grupos, columna, fijo, hasta_fila=fila_titulos)

    columna = 4
    donde = {}
    for codigo, etiqueta in Dimension.choices:
        suyas = [a for a in actividades if a.dimension == codigo]
        if not suyas:
            continue
        _titulo(hoja, fila_grupos, columna, f'{etiqueta} (sobre {MAXIMO[codigo]:.0f})', FONDO_MATERIA,
                hasta_columna=columna + len(suyas))
        for actividad in suyas:
            _titulo(hoja, fila_titulos, columna, actividad.etiqueta, FONDO_MATERIA)
            donde[actividad.id] = columna
            columna += 1
        _titulo(hoja, fila_titulos, columna, f'Prom. {etiqueta.lower()}', FONDO_MATERIA)
        donde[f'promedio_{codigo}'] = columna
        columna += 1

    _titulo(hoja, fila_grupos, columna, 'Cierre del trimestre', FONDO_RESUMEN, hasta_columna=columna + 2)
    for desplazamiento, etiqueta in enumerate([
        f'Autoevaluación (sobre {MAXIMO_AUTOEVALUACION:.0f})',
        f'Extracurricular (hasta {MAXIMO_EXTRACURRICULAR:.0f})',
        'Nota del trimestre',
    ]):
        _titulo(hoja, fila_titulos, columna + desplazamiento, etiqueta, FONDO_RESUMEN)
    ultima = columna + 2

    for indice, matricula in enumerate(matriculas):
        r = primera + indice
        estudiante = matricula.estudiante
        nota = notas.get(estudiante.pk)
        hoja.cell(row=r, column=1, value=indice + 1).alignment = CENTRADO
        hoja.cell(row=r, column=2, value=estudiante.rude).alignment = CENTRADO
        hoja.cell(row=r, column=3, value=estudiante.usuario.get_full_name()).alignment = IZQUIERDA

        for actividad in actividades:
            valor = puntajes.get((actividad.id, estudiante.pk))
            celda = hoja.cell(row=r, column=donde[actividad.id])
            celda.alignment = CENTRADO
            if valor is None:
                celda.value = '—'
                celda.font = APAGADO
            else:
                celda.value = float(valor)

        for codigo, _etiqueta in Dimension.choices:
            columna_promedio = donde.get(f'promedio_{codigo}')
            if columna_promedio is None:
                continue
            valor = getattr(nota, codigo.lower(), None) if nota else None
            celda = hoja.cell(row=r, column=columna_promedio)
            celda.alignment = CENTRADO
            celda.font = NEGRITA
            celda.value = float(valor) if valor is not None else '—'

        for desplazamiento, parte in enumerate(['autoevaluacion', 'extracurricular']):
            valor = getattr(nota, parte, None) if nota else None
            celda = hoja.cell(row=r, column=ultima - 2 + desplazamiento)
            celda.alignment = CENTRADO
            celda.value = float(valor) if valor is not None else '—'
        _nota(hoja.cell(row=r, column=ultima), float(nota.nota) if nota and nota.nota else None)

    hoja.column_dimensions['A'].width = 5
    hoja.column_dimensions['B'].width = 18
    hoja.column_dimensions['C'].width = 34
    for c in range(4, ultima + 1):
        hoja.column_dimensions[get_column_letter(c)].width = 13
    hoja.row_dimensions[fila_titulos].height = 30
    hoja.freeze_panes = hoja.cell(row=primera, column=4)
    return libro
