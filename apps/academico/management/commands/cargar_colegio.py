"""Carga el colegio real desde el Excel de distribución de cursos.

El Excel trae dos hojas y se necesitan las dos:

  Hoja1  — quién dicta qué materia en qué curso (con apodos: "SULMA", "BETTY B.")
  Hoja2  — dónde está cada docente en cada hora de la semana (nombre y apellido)

Ninguna basta sola: la Hoja2 dice que Sulma está en 4°C el martes a primera hora,
pero no cuál de sus tres materias toca; eso lo dice la Hoja1.

El comando es idempotente: correrlo dos veces deja la base igual que una.
"""
import collections
import re
import unicodedata
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.academico.models import (
    AsignacionDocente, Curso, DiaCapacitacion, Estudiante, Gestion, Horario,
    Materia, Matricula, Periodo,
)
from apps.accounts.models import Usuario

DIAS = ['LUN', 'MAR', 'MIE', 'JUE', 'VIE', 'SAB']

# Horas de clase del Excel: siete columnas por día.
PERIODOS_POR_DIA = 7

# El recreo parte el día entre la cuarta y la quinta hora. Ocupa su propia fila
# de la grilla, así que la quinta hora de clase vive en el orden 5, no en el 4.
ORDEN_DEL_RECREO = 4


def orden_del_periodo(hora_de_clase):
    """De la hora del Excel (0-6) a la fila de la grilla, salteando el recreo."""
    return hora_de_clase if hora_de_clase < ORDEN_DEL_RECREO else hora_de_clase + 1

# Un docente con cinco o más horas en el mismo curso el mismo día no está dando
# clase: es la capacitación de la mañana, que va por fuera de este horario de
# tarde. El bloque legítimo más largo del colegio es de cuatro horas.
HORAS_QUE_HACEN_TALLER = 5

# Nombres completos tal como los entregó el colegio. La Hoja2 solo trae nombre y
# primer apellido, que es lo que forma el usuario; esto es para mostrarlos.
NOMBRES_COMPLETOS = [
    'Fausto Enrique Vargas Herrera', 'Wilfredo Eduardo Bitre Hinojosa',
    'Miriam Cussi Ancasi', 'Jose Luis Rodriguez Yampasi', 'Mery Alvarez Mamani',
    'Betty Rosario Borda Soto', 'Gaby Judith Lopez Apaza',
    'Gerardo Ramiro Aguilar Fernandez', 'Ada Efigenia Campos Quiroz',
    'Lidia Angelica Castillo Mamani', 'Jose Elias Condorena Ticona',
    'Elizabeth Chuquimia Sequeiros', 'Gustavo Gutierrez Quispe',
    'Mariela Chura Roque', 'Ramiro Eduardo Barrientos Valdez',
    'Gimena Fernandez Laura', 'Jorge Luis Vitre Apaza',
    'Limbert Remberto Soliz Choque', 'Eduardo Rojas Quispe',
    'Sonia Ruth Colque Condori', 'Aydee Erika Plata Conde',
    'Elsa Nieves Payllo Zarzuri', 'Juan Carlos Chipana Quispe',
    'Susy Francisca Chuquimia Tinta', 'Bertha Quispe Lucana',
    'Sulma Nelly Yujra Loza', 'Hernan Choquehuanca Loza',
    'Nelly Nancy Conde Limachi', 'Veronica Lucia Alcon Merlo',
    'Isidro Carlos Coaquira Hilari', 'Nelly Felisa Blanco Rodriguez',
    'Margarita Janeth Merlo Paredes', 'Jesusa Bonifacia Quispe Chipana',
    'Margot Ninoska Choque Laure', 'Juan Carlos Perez Guzman',
    'Ruben Freddy Ledezma Vera', 'Celia Ramos Choquehuanca',
]


def sin_tildes(texto):
    return unicodedata.normalize('NFD', texto or '').encode('ascii', 'ignore').decode()


def normalizar(texto):
    """Para comparar: sin tildes, sin puntos, en mayúsculas, un solo espacio."""
    return re.sub(r'\s+', ' ', sin_tildes(texto).replace('.', ' ')).strip().upper()


def nombre_de_usuario(nombre_y_apellido):
    """De "Veronica Alcon" sale "veronicaalcon"."""
    return ''.join(c for c in sin_tildes(nombre_y_apellido).lower() if c.isalnum())


class Command(BaseCommand):
    help = 'Carga docentes, materias, asignaciones y horario desde el Excel del colegio.'

    def add_arguments(self, parser):
        parser.add_argument('excel', help='Ruta del archivo distribucion cursos.xlsx')
        parser.add_argument(
            '--borrar-demo', action='store_true',
            help='Borra antes los estudiantes y profesores de prueba con todo lo suyo.',
        )

    @transaction.atomic
    def handle(self, *args, **opciones):
        ruta = Path(opciones['excel'])
        if not ruta.exists():
            raise CommandError(f'No encuentro el archivo: {ruta}')

        gestion = Gestion.objects.filter(activa=True).order_by('-anio').first()
        if gestion is None:
            raise CommandError('No hay ninguna gestión activa. Créala antes de cargar.')

        hoja1, hoja2 = self._leer_excel(ruta)
        self.stdout.write(
            f'Excel: {len(hoja2)} docentes, {len(hoja1)} asignaciones curso-materia'
        )

        if opciones['borrar_demo']:
            self._borrar_demo()

        materias = self._crear_materias(hoja1)
        docentes = self._crear_docentes(hoja2, materias)
        apodos = self._resolver_apodos(hoja1, hoja2, docentes)
        cursos = self._cursos(gestion)
        periodos = self._crear_periodos(cursos)
        asignaciones = self._crear_asignaciones(hoja1, apodos, materias, cursos)
        self._crear_horario(hoja2, hoja1, apodos, asignaciones, periodos, cursos, docentes)

    # ────────────────────────────── lectura ──────────────────────────────

    def _leer_excel(self, ruta):
        from openpyxl import load_workbook

        libro = load_workbook(ruta, data_only=True)
        if 'Hoja1' not in libro.sheetnames or 'Hoja2 (2)' not in libro.sheetnames:
            raise CommandError(
                f'El Excel debe traer "Hoja1" y "Hoja2 (2)". Trae: {libro.sheetnames}'
            )

        # ── Hoja1: la rejilla materia x curso ──
        h1 = list(libro['Hoja1'].iter_rows(min_row=1, max_row=40, max_col=22, values_only=True))
        fila_grado = [('' if c is None else str(c).strip()) for c in h1[2]]
        fila_par = [('' if c is None else str(c).strip()) for c in h1[3]]
        columnas, grado = {}, None
        for i in range(1, 20):
            if i < len(fila_grado) and fila_grado[i]:
                grado = fila_grado[i][0]
            if grado and i < len(fila_par) and fila_par[i] in ('A', 'B', 'C'):
                columnas[i] = f'{grado}{fila_par[i]}'

        asignaciones = {}
        for fila in h1[4:]:
            materia = ('' if fila[0] is None else str(fila[0]).strip())
            if not materia:
                continue
            for i, curso in columnas.items():
                if i < len(fila) and fila[i]:
                    asignaciones[(curso, materia)] = str(fila[i]).strip()

        # ── Hoja2: el horario por docente ──
        docentes = {}
        for fila in libro['Hoja2 (2)'].iter_rows(min_row=4, max_row=40, values_only=True):
            nombre = (fila[1] or '').strip()
            if not nombre:
                continue
            materias = [str(fila[i]).strip() for i in (2, 3, 4) if i < len(fila) and fila[i]]
            celdas = {}
            for d, dia in enumerate(DIAS):
                for p in range(PERIODOS_POR_DIA):
                    col = 5 + d * PERIODOS_POR_DIA + p
                    if col < len(fila) and fila[col]:
                        celdas[(dia, p)] = str(fila[col]).strip().upper().replace(' ', '')
            docentes[nombre] = {'materias': materias, 'celdas': celdas}

        return asignaciones, docentes

    # ────────────────────────────── borrado ──────────────────────────────

    def _borrar_demo(self):
        """Se va todo lo inventado. Quedan admin, regente, la gestión y los cursos."""
        from apps.asistencia.models import Asistencia
        from apps.calificaciones.models import Nota
        from apps.comunicaciones.models import Citacion
        from apps.mensajeria.models import Conversacion
        from apps.notificaciones.models import Notificacion
        from apps.tareas.models import EntregaTarea, Tarea

        usuarios_fuera = list(
            Usuario.objects.filter(
                rol__in=[Usuario.Rol.ESTUDIANTE, Usuario.Rol.PADRE, Usuario.Rol.PROFESOR]
            ).values_list('pk', flat=True)
        )
        contados = {
            'notas': Nota.objects.count(),
            'asistencias': Asistencia.objects.count(),
            'entregas de tareas': EntregaTarea.objects.count(),
            'tareas': Tarea.objects.count(),
            'citaciones': Citacion.objects.count(),
            'conversaciones': Conversacion.objects.count(),
            'horarios': Horario.objects.count(),
            'periodos': Periodo.objects.count(),
            'asignaciones': AsignacionDocente.objects.count(),
            'matriculas': Matricula.objects.count(),
            'estudiantes': Estudiante.objects.count(),
            'materias': Materia.objects.count(),
            'usuarios': len(usuarios_fuera),
        }

        Nota.objects.all().delete()
        Asistencia.objects.all().delete()
        EntregaTarea.objects.all().delete()
        Tarea.objects.all().delete()
        Citacion.objects.all().delete()
        Conversacion.objects.all().delete()
        Notificacion.objects.filter(usuario_id__in=usuarios_fuera).delete()
        Horario.objects.all().delete()
        Periodo.objects.all().delete()
        AsignacionDocente.objects.all().delete()
        DiaCapacitacion.objects.all().delete()
        Matricula.objects.all().delete()
        Estudiante.objects.all().delete()
        Usuario.objects.filter(pk__in=usuarios_fuera).delete()
        # Las materias se van con todo lo demas y se vuelven a crear desde el
        # Excel; si no, quedan las dos inventadas ensuciando los desplegables.
        Materia.objects.all().delete()

        self.stdout.write(self.style.WARNING('Datos de prueba borrados:'))
        for que, cuantos in contados.items():
            if cuantos:
                self.stdout.write(f'  {que}: {cuantos}')

    # ────────────────────────────── creación ──────────────────────────────

    def _crear_materias(self, hoja1):
        materias = {}
        for nombre in sorted({materia for (_curso, materia) in hoja1}):
            materia, _ = Materia.objects.get_or_create(nombre=nombre)
            materias[normalizar(nombre)] = materia
        self.stdout.write(f'Materias: {len(materias)}')
        return materias

    def _crear_docentes(self, hoja2, materias):
        # La Hoja2 trae "VERONICA ALCON"; la fotografía, "Veronica Lucia Alcon
        # Merlo". Se emparejan por nombre de pila y primer apellido.
        completos = {}
        for completo in NOMBRES_COMPLETOS:
            partes = completo.split()
            if len(partes) >= 2:
                completos[normalizar(f'{partes[0]} {partes[-2]}')] = completo

        docentes = {}
        for nombre_corto, datos in hoja2.items():
            usuario_id = nombre_de_usuario(nombre_corto)
            completo = completos.get(normalizar(nombre_corto), nombre_corto.title())
            partes = completo.split()

            usuario, _ = Usuario.objects.get_or_create(
                username=usuario_id, defaults={'rol': Usuario.Rol.PROFESOR},
            )
            usuario.rol = Usuario.Rol.PROFESOR
            usuario.is_active = True
            if len(partes) >= 3:
                usuario.first_name = ' '.join(partes[:-2])
                usuario.last_name = ' '.join(partes[-2:])
            else:
                usuario.first_name = partes[0]
                usuario.last_name = ' '.join(partes[1:])
            # Sin carnet todavía: la contraseña es el mismo usuario. Se cambia
            # cuando el colegio entregue los CI.
            usuario.set_password(usuario_id)
            usuario.save()
            docentes[nombre_corto] = usuario

        con_varias = sum(1 for d in hoja2.values() if len(d['materias']) > 1)
        self.stdout.write(
            f'Docentes: {len(docentes)} ({con_varias} con más de una materia). '
            'Usuario y contraseña iguales, sin tildes ni espacios.'
        )
        return docentes

    def _resolver_apodos(self, hoja1, hoja2, docentes):
        """De "SULMA" o "BETTY B." al docente de la Hoja2 que corresponde."""
        por_apodo = collections.defaultdict(set)
        for (curso, materia), apodo in hoja1.items():
            por_apodo[apodo].add((curso, materia))

        resueltos, dudosos = {}, []
        for apodo, pares in por_apodo.items():
            candidatos = []
            for nombre, datos in hoja2.items():
                suyas = {normalizar(m) for m in datos['materias']}
                sus_cursos = set(datos['celdas'].values())
                if all(normalizar(m) in suyas and c in sus_cursos for (c, m) in pares):
                    candidatos.append(nombre)
            if len(candidatos) > 1:
                # Desempate por nombre: "CARLOS C" es principio de "CARLOS COAQUIRA".
                inicio = normalizar(apodo)
                por_nombre = [n for n in candidatos if normalizar(n).startswith(inicio)]
                candidatos = por_nombre or candidatos
            if len(candidatos) == 1:
                resueltos[apodo] = docentes[candidatos[0]]
            else:
                dudosos.append((apodo, candidatos))

        if dudosos:
            detalle = '; '.join(f'"{a}" -> {c or "nadie"}' for a, c in dudosos)
            raise CommandError(f'No pude identificar a estos docentes de la Hoja1: {detalle}')
        self.stdout.write(f'Apodos de la Hoja1 resueltos: {len(resueltos)}')
        return resueltos

    def _cursos(self, gestion):
        cursos = {
            f'{curso.grado}{curso.paralelo}': curso
            for curso in Curso.objects.filter(gestion=gestion)
        }
        self.stdout.write(f'Cursos de la gestión {gestion.anio}: {len(cursos)}')
        return cursos

    def _crear_periodos(self, cursos):
        """Siete horas de clase más el recreo, que va entre la cuarta y la quinta."""
        periodos, filas = {}, 0
        for clave, curso in cursos.items():
            for hora in range(PERIODOS_POR_DIA):
                orden = orden_del_periodo(hora)
                periodo, _ = Periodo.objects.update_or_create(
                    curso=curso, orden=orden,
                    defaults={'etiqueta': f'{hora + 1}°', 'es_recreo': False},
                )
                periodos[(clave, hora)] = periodo
                filas += 1
            Periodo.objects.update_or_create(
                curso=curso, orden=ORDEN_DEL_RECREO,
                defaults={'etiqueta': 'Recreo', 'es_recreo': True},
            )
            filas += 1
        self.stdout.write(
            f'Períodos: {filas} ({PERIODOS_POR_DIA} horas de clase y el recreo por curso)'
        )
        return periodos

    def _crear_asignaciones(self, hoja1, apodos, materias, cursos):
        asignaciones, sin_curso = {}, set()
        for (clave_curso, nombre_materia), apodo in hoja1.items():
            curso = cursos.get(clave_curso)
            if curso is None:
                sin_curso.add(clave_curso)
                continue
            asignacion, _ = AsignacionDocente.objects.update_or_create(
                curso=curso, materia=materias[normalizar(nombre_materia)],
                defaults={'profesor': apodos[apodo]},
            )
            asignaciones[(clave_curso, normalizar(nombre_materia))] = asignacion
        if sin_curso:
            self.stdout.write(self.style.WARNING(
                f'  Cursos del Excel que no existen en la base: {sorted(sin_curso)}'
            ))
        self.stdout.write(f'Asignaciones curso-materia: {len(asignaciones)}')
        return asignaciones

    def _crear_horario(self, hoja2, hoja1, apodos, asignaciones, periodos, cursos, docentes):
        # Un mismo docente aparece en la Hoja1 con varios apodos segun la materia:
        # Betty Borda es "BETTY" en Comunicacion y "BETTY B." en Sociales.
        sus_apodos = collections.defaultdict(set)
        for apodo, usuario in apodos.items():
            sus_apodos[usuario.pk].add(apodo)

        # Bloques de día completo: son la capacitación de la mañana, no van aquí.
        bloques = collections.Counter()
        for nombre, datos in hoja2.items():
            for (dia, _orden), curso in datos['celdas'].items():
                bloques[(nombre, curso, dia)] += 1
        talleres = {clave for clave, horas in bloques.items() if horas >= HORAS_QUE_HACEN_TALLER}

        creados, fuera = 0, 0
        ambiguas, sin_materia, choques = [], [], []
        ocupadas = {}

        for nombre, datos in hoja2.items():
            mios = sus_apodos.get(docentes[nombre].pk, set())
            for (dia, orden), clave_curso in sorted(datos['celdas'].items()):
                if (nombre, clave_curso, dia) in talleres:
                    fuera += 1
                    continue
                if clave_curso not in cursos:
                    continue

                posibles = [m for m in datos['materias'] if hoja1.get((clave_curso, m)) in mios]
                if not posibles:
                    sin_materia.append((nombre, clave_curso, dia, orden + 1))
                    continue
                if len(posibles) > 1:
                    ambiguas.append((nombre, clave_curso, dia, orden + 1, posibles))

                asignacion = asignaciones.get((clave_curso, normalizar(posibles[0])))
                if asignacion is None:
                    sin_materia.append((nombre, clave_curso, dia, orden + 1))
                    continue

                celda = (clave_curso, dia, orden)
                if celda in ocupadas:
                    choques.append((clave_curso, dia, orden + 1, ocupadas[celda], nombre))
                    continue

                Horario.objects.update_or_create(
                    periodo=periodos[(clave_curso, orden)], dia=dia,
                    defaults={'asignacion': asignacion},
                )
                ocupadas[celda] = nombre
                creados += 1

        self.stdout.write(f'Horas de clase cargadas: {creados}')
        self.stdout.write(f'  fuera del horario por ser taller de la mañana: {fuera}')

        # Los días de taller pasan a ser días de capacitación, para que la
        # administración pueda tomarles asistencia.
        dias_taller = {(curso, dia) for (_nombre, curso, dia) in talleres}
        for clave_curso, dia in sorted(dias_taller):
            if clave_curso in cursos:
                DiaCapacitacion.objects.get_or_create(curso=cursos[clave_curso], dia=dia)
        self.stdout.write(f'Días de capacitación: {len(dias_taller)}')

        self._avisar(choques, ambiguas, sin_materia)
        self.stdout.write(self.style.SUCCESS('\nListo.'))

    def _avisar(self, choques, ambiguas, sin_materia):
        if choques:
            self.stdout.write(self.style.WARNING(
                f'\n{len(choques)} choque(s) del propio Excel: dos docentes a la misma hora '
                'en el mismo curso. Se cargó el primero; corrige en "Armar horarios":'))
            for curso, dia, periodo, primero, segundo in choques:
                self.stdout.write(f'  {curso} {dia} período {periodo}: {primero} (cargado) / {segundo}')

        if ambiguas:
            self.stdout.write(self.style.WARNING(
                f'\n{len(ambiguas)} hora(s) donde el docente dicta varias materias al mismo '
                'curso y el Excel no dice cuál toca. Se cargó la primera:'))
            for nombre, curso, dia, periodo, posibles in ambiguas[:15]:
                self.stdout.write(f'  {nombre} · {curso} {dia} período {periodo}: {" / ".join(posibles)}')
            if len(ambiguas) > 15:
                self.stdout.write(f'  ... y {len(ambiguas) - 15} más')

        if sin_materia:
            self.stdout.write(self.style.WARNING(
                f'\n{len(sin_materia)} hora(s) sin materia deducible, no se cargaron:'))
            for nombre, curso, dia, periodo in sin_materia[:10]:
                self.stdout.write(f'  {nombre} · {curso} {dia} período {periodo}')
            if len(sin_materia) > 10:
                self.stdout.write(f'  ... y {len(sin_materia) - 10} más')
