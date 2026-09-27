"""Planilla de asistencia de un curso, en Excel.

**Una hoja por mes.** Un año escolar entero en una sola hoja son unas doscientas
columnas: no se lee en pantalla ni se imprime. Separado por meses, cada hoja es
justo el registro mensual que se archiva.

La primera hoja resume el año, para no perder de vista el acumulado al repartir
las fechas entre hojas.
"""
import datetime
import re

from django.utils import formats, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Asistencia

FILA_TITULOS = 4

# Fondos suaves para leer la planilla de un vistazo, sin ir palabra por palabra.
# Son los colores del sistema, aclarados.
RELLENO_ESTADO = {
    Asistencia.Estado.PRESENTE: 'D6F0E4',
    Asistencia.Estado.FALTA: 'F8D8D4',
    Asistencia.Estado.ATRASO: 'FCE9CE',
    Asistencia.Estado.LICENCIA: 'DEE5F7',
}

CENTRADO = Alignment(horizontal='center', vertical='center')
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
APAGADO = Font(italic=True, color='666666')
FONDO_TITULOS = PatternFill('solid', fgColor='EDF6F4')


def nombre_de_hoja(texto):
    """Excel no admite \\ / * ? : [ ] en el nombre de la hoja, y corta en 31."""
    return re.sub(r'[\\/*?:\[\]]', ' ', texto)[:31] or 'Asistencia'


def meses_de(fechas):
    """Agrupa las fechas por mes conservando el orden cronológico."""
    meses = {}
    for fecha in sorted(fechas):
        meses.setdefault((fecha.year, fecha.month), []).append(fecha)
    return meses


def nombre_del_mes(anio, mes):
    return f'{formats.date_format(datetime.date(anio, mes, 1), "F").capitalize()} {anio}'


def _encabezado(hoja, curso, subtitulo):
    hoja['A1'] = f'Asistencia de capacitación · {curso}'
    hoja['A1'].font = TITULO
    hoja['A2'] = subtitulo
    hoja['A2'].font = APAGADO


def _titulos(hoja, textos):
    for columna, texto in enumerate(textos, start=1):
        celda = hoja.cell(row=FILA_TITULOS, column=columna, value=texto)
        celda.font = NEGRITA
        celda.alignment = CENTRADO
        celda.fill = FONDO_TITULOS


def _identificar(hoja, fila, estudiante):
    hoja.cell(row=fila, column=1, value=estudiante.usuario.get_full_name())
    hoja.cell(row=fila, column=2, value=estudiante.rude).alignment = CENTRADO


def _anchos(hoja, columnas_de_datos, ancho=11):
    hoja.column_dimensions['A'].width = 30
    hoja.column_dimensions['B'].width = 12
    for columna in range(3, 3 + columnas_de_datos):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    # El nombre y el RUDE quedan a la vista al desplazarse por las fechas.
    hoja.freeze_panes = hoja.cell(row=FILA_TITULOS + 1, column=3)


def _cuenta(registros):
    """Presentes, faltas, asistidos y registrados de un conjunto de registros."""
    presentes = sum(1 for r in registros if r.estado == Asistencia.Estado.PRESENTE)
    faltas = sum(1 for r in registros if r.estado == Asistencia.Estado.FALTA)
    asistidos = sum(1 for r in registros if r.asistio)
    return presentes, faltas, asistidos, len(registros)


def _totales(hoja, fila, columna, presentes, faltas, asistidos, registrados):
    hoja.cell(row=fila, column=columna, value=presentes).alignment = CENTRADO
    hoja.cell(row=fila, column=columna + 1, value=faltas).alignment = CENTRADO
    porcentaje = hoja.cell(
        row=fila, column=columna + 2,
        value=(asistidos / registrados) if registrados else None,
    )
    porcentaje.number_format = '0%'
    porcentaje.alignment = CENTRADO


def _hoja_del_mes(libro, curso, matriculas, fechas, por_estudiante, anio, mes):
    """Una hoja con los días de ese mes: es el registro mensual de siempre."""
    hoja = libro.create_sheet(nombre_de_hoja(nombre_del_mes(anio, mes)))
    _encabezado(hoja, curso, nombre_del_mes(anio, mes))

    _titulos(
        hoja,
        ['Estudiante', 'RUDE']
        + [formats.date_format(f, 'D d') for f in fechas]
        + ['Presentes', 'Faltas', 'Asistencia'],
    )

    for indice, matricula in enumerate(matriculas):
        fila = FILA_TITULOS + 1 + indice
        del_estudiante = por_estudiante.get(matricula.estudiante_id, {})
        _identificar(hoja, fila, matricula.estudiante)

        del_mes = []
        for desplazamiento, fecha in enumerate(fechas):
            registro = del_estudiante.get(fecha)
            celda = hoja.cell(row=fila, column=3 + desplazamiento)
            celda.alignment = CENTRADO
            if registro is None:
                celda.value = '—'
                continue
            celda.value = registro.get_estado_display()
            relleno = RELLENO_ESTADO.get(registro.estado)
            if relleno:
                celda.fill = PatternFill('solid', fgColor=relleno)
            del_mes.append(registro)

        _totales(hoja, fila, 3 + len(fechas), *_cuenta(del_mes))

    # Cuántos vinieron cada día: es lo primero que se mira al abrir la hoja.
    if matriculas:
        fila = FILA_TITULOS + 1 + len(matriculas)
        hoja.cell(row=fila, column=1, value='Presentes del día').font = NEGRITA
        for desplazamiento, fecha in enumerate(fechas):
            cuantos = sum(
                1 for m in matriculas
                for r in [por_estudiante.get(m.estudiante_id, {}).get(fecha)]
                if r is not None and r.asistio
            )
            celda = hoja.cell(row=fila, column=3 + desplazamiento, value=cuantos)
            celda.font = NEGRITA
            celda.alignment = CENTRADO

    _anchos(hoja, len(fechas) + 3)
    return hoja


def _hoja_resumen(libro, curso, gestion, matriculas, meses, por_estudiante):
    """El acumulado del año: faltas de cada mes y total, en una sola mirada."""
    hoja = libro.create_sheet('Resumen', 0)
    _encabezado(hoja, curso, (
        f'Gestión {gestion.anio} · '
        f'generado el {formats.date_format(timezone.localdate(), "j \\d\\e F \\d\\e Y")}'
    ))

    claves = list(meses)
    _titulos(
        hoja,
        ['Estudiante', 'RUDE']
        + [f'Faltas · {nombre_del_mes(anio, mes)}' for anio, mes in claves]
        + ['Presentes', 'Faltas', 'Asistencia'],
    )

    for indice, matricula in enumerate(matriculas):
        fila = FILA_TITULOS + 1 + indice
        del_estudiante = por_estudiante.get(matricula.estudiante_id, {})
        _identificar(hoja, fila, matricula.estudiante)

        for desplazamiento, (anio, mes) in enumerate(claves):
            faltas = sum(
                1 for fecha in meses[(anio, mes)]
                for r in [del_estudiante.get(fecha)]
                if r is not None and r.estado == Asistencia.Estado.FALTA
            )
            celda = hoja.cell(row=fila, column=3 + desplazamiento, value=faltas)
            celda.alignment = CENTRADO
            if faltas:
                celda.fill = PatternFill('solid', fgColor=RELLENO_ESTADO[Asistencia.Estado.FALTA])

        _totales(hoja, fila, 3 + len(claves), *_cuenta(list(del_estudiante.values())))

    if matriculas and claves:
        fila = FILA_TITULOS + 1 + len(matriculas)
        hoja.cell(row=fila, column=1, value='Faltas del curso').font = NEGRITA
        for desplazamiento, (anio, mes) in enumerate(claves):
            total = sum(
                1 for m in matriculas
                for fecha in meses[(anio, mes)]
                for r in [por_estudiante.get(m.estudiante_id, {}).get(fecha)]
                if r is not None and r.estado == Asistencia.Estado.FALTA
            )
            celda = hoja.cell(row=fila, column=3 + desplazamiento, value=total)
            celda.font = NEGRITA
            celda.alignment = CENTRADO

    if not claves:
        hoja.cell(
            row=FILA_TITULOS + len(matriculas) + 2, column=1,
            value='Todavía no hay asistencia registrada para este curso.',
        ).font = APAGADO

    _anchos(hoja, len(claves) + 3, ancho=18 if claves else 11)
    return hoja


def planilla_de_capacitacion(curso, gestion, matriculas, fechas, por_estudiante):
    """Libro completo: hoja de resumen y después una hoja por mes."""
    libro = Workbook()
    libro.remove(libro.active)          # las hojas se crean a medida

    meses = meses_de(fechas)
    _hoja_resumen(libro, curso, gestion, matriculas, meses, por_estudiante)
    for (anio, mes), del_mes in meses.items():
        _hoja_del_mes(libro, curso, matriculas, del_mes, por_estudiante, anio, mes)
    return libro
