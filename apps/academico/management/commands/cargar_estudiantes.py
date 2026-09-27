"""Carga los estudiantes desde las listas del SIE («Estudiantes inscritos por curso»).

Cada Excel trae una hoja por curso, con su grado y paralelo en la cabecera, más
una hoja «Consolidado» y otra «Resumen» que repiten lo mismo y se saltean.
Un curso puede venir en dos archivos (1ro A viene así): se toma una vez.

Reglas del colegio:
  - la columna «Obs. Bono Familia» no se lee;
  - quien figura como RETIRADO TRASLADO no se carga, y si ya estaba cargado su
    matrícula se desactiva (no se borra: conserva su historial).

El comando es idempotente: correrlo dos veces deja la base igual que una.
"""
import collections
import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.academico.models import Curso, Estudiante, Gestion, Matricula
from apps.accounts.models import Usuario

GRADOS = {'PRIMERO': 1, 'SEGUNDO': 2, 'TERCERO': 3, 'CUARTO': 4, 'QUINTO': 5, 'SEXTO': 6}

RETIRADO = 'RETIRADO TRASLADO'

# Las columnas que se usan. «Obs. Bono Familia», provincia y localidad no.
COLUMNAS = {
    'rude': 'Código RUDE', 'carnet': 'Carnet', 'nombre': 'Nombre Completo',
    'genero': 'Género', 'nacimiento': 'Fecha Nacimiento', 'pais': 'País',
    'departamento': 'Departamento', 'matricula': 'Matrícula',
}

NACIONALIDADES = {'BOLIVIA': 'Boliviana', 'ARGENTINA': 'Argentina', 'BRASIL': 'Brasileña', 'PERU': 'Peruana', 'CHILE': 'Chilena'}

DEPARTAMENTOS = {
    'LA PAZ': Estudiante.Departamento.LA_PAZ, 'COCHABAMBA': Estudiante.Departamento.COCHABAMBA,
    'SANTA CRUZ': Estudiante.Departamento.SANTA_CRUZ, 'ORURO': Estudiante.Departamento.ORURO,
    'POTOSI': Estudiante.Departamento.POTOSI, 'POTOSÍ': Estudiante.Departamento.POTOSI,
    'CHUQUISACA': Estudiante.Departamento.CHUQUISACA, 'TARIJA': Estudiante.Departamento.TARIJA,
    'BENI': Estudiante.Departamento.BENI, 'PANDO': Estudiante.Departamento.PANDO,
}

# La lista trae «APELLIDOS NOMBRES» todo junto y lo normal es llevar dos
# apellidos. Estas personas, con solo tres palabras, llevan uno: se comprobó
# cruzando con los apellidos y nombres del resto de la lista.
UN_SOLO_APELLIDO = {
    'CALDERON JHONATAN RICARDO', 'FLORES IGNACIO BRANDON', 'GARCIA KENDRA SHERLYN',
    'GOMEZ KARINA SHEYLA', 'OCHOA MILENKA AURORA', 'TARQUI RAMIRO FELIX',
    'UCEDO MIGUEL ANGEL',
}

PARTICULAS = {'DE', 'DEL', 'LA', 'LAS', 'LOS', 'Y'}


def texto(valor):
    return '' if valor is None else str(valor).strip()


def capitalizar(palabras):
    """«ANGEL DE JESUS» → «Angel de Jesus»."""
    return ' '.join(
        p.lower() if i and p in PARTICULAS else p.capitalize()
        for i, p in enumerate(palabras)
    )


def separar_nombre(nombre_completo):
    """De «ACHO ESPINOZA JHASMANY RAUL» a («Jhasmany Raul», «Acho Espinoza»)."""
    palabras = nombre_completo.upper().split()
    cuantos = 1 if ' '.join(palabras) in UN_SOLO_APELLIDO or len(palabras) < 3 else 2
    return capitalizar(palabras[cuantos:]), capitalizar(palabras[:cuantos])


class Command(BaseCommand):
    help = 'Carga estudiantes y matrículas desde las listas del SIE por curso.'

    def add_arguments(self, parser):
        parser.add_argument('excel', nargs='+', help='Uno o más Excel de estudiantes por curso.')

    @transaction.atomic
    def handle(self, *args, **opciones):
        rutas = [Path(r) for r in opciones['excel']]
        faltan = [str(r) for r in rutas if not r.exists()]
        if faltan:
            raise CommandError(f'No encuentro: {faltan}')

        gestion = Gestion.objects.filter(activa=True).order_by('-anio').first()
        if gestion is None:
            raise CommandError('No hay ninguna gestión activa. Créala antes de cargar.')

        filas = self._leer(rutas)
        cursos = self._cursos(gestion, filas)
        vigentes = [f for f in filas.values() if f['matricula'] != RETIRADO]
        retirados = [f for f in filas.values() if f['matricula'] == RETIRADO]
        self.stdout.write(
            f'Listas: {len(filas)} estudiantes distintos · '
            f'{len(retirados)} retirados que no se cargan · {len(vigentes)} a cargar'
        )

        self._revisar_choques(vigentes)
        creados, actualizados = self._guardar(vigentes, cursos, gestion)
        desactivados = self._desactivar_retirados(retirados, gestion)
        self._resumen(vigentes, retirados, creados, actualizados, desactivados, cursos)

    # ────────────────────────────── lectura ──────────────────────────────

    def _leer(self, rutas):
        from openpyxl import load_workbook

        filas, diferencias = {}, []
        for ruta in rutas:
            libro = load_workbook(ruta, read_only=True, data_only=True)
            for hoja in libro.worksheets:
                if hoja.title.strip().upper() in ('CONSOLIDADO', 'RESUMEN'):
                    continue
                for fila in self._leer_hoja(hoja, ruta):
                    anterior = filas.get(fila['rude'])
                    if anterior is None:
                        filas[fila['rude']] = fila
                    elif self._datos(anterior) != self._datos(fila):
                        diferencias.append(
                            f"  RUDE {fila['rude']}: {anterior['origen']} dice "
                            f"{self._datos(anterior)}, {fila['origen']} dice {self._datos(fila)}"
                        )
            libro.close()

        if diferencias:
            raise CommandError(
                'El mismo RUDE aparece con datos distintos; corrige las listas:\n'
                + '\n'.join(diferencias)
            )
        if not filas:
            raise CommandError('No encontré ninguna hoja de curso con estudiantes.')
        return filas

    def _leer_hoja(self, hoja, ruta):
        grado = paralelo = None
        titulos = None
        for valores in hoja.iter_rows(values_only=True):
            if not valores or all(v is None for v in valores):
                continue
            primera = texto(valores[0])
            if titulos is None:
                if primera == 'Grado:':
                    grado = GRADOS.get(texto(valores[1]).upper())
                elif primera == 'Paralelo:':
                    paralelo = texto(valores[1]).upper()
                elif COLUMNAS['rude'] in [texto(v) for v in valores]:
                    titulos = {texto(v): i for i, v in enumerate(valores)}
                    faltan = [c for c in COLUMNAS.values() if c not in titulos]
                    if faltan:
                        raise CommandError(f'{ruta.name} / {hoja.title}: faltan las columnas {faltan}')
                    if grado is None or not paralelo:
                        raise CommandError(
                            f'{ruta.name} / {hoja.title}: no encuentro «Grado:» y «Paralelo:» en la cabecera'
                        )
                continue

            rude = texto(valores[titulos[COLUMNAS['rude']]])
            if not rude:
                continue
            fila = {
                clave: valores[titulos[columna]] if titulos[columna] < len(valores) else None
                for clave, columna in COLUMNAS.items()
            }
            fila.update(
                rude=rude.upper(), carnet=texto(fila['carnet']).upper(),
                nombre=' '.join(texto(fila['nombre']).split()),
                matricula=texto(fila['matricula']).upper(),
                curso=(grado, paralelo), origen=f'{ruta.name} / {hoja.title}',
            )
            if not fila['carnet']:
                raise CommandError(f"{fila['origen']}: {fila['nombre']} no tiene carnet")
            yield fila

    @staticmethod
    def _datos(fila):
        """Lo que tiene que coincidir si un estudiante viene en dos listas."""
        return (fila['carnet'], fila['nombre'], fila['curso'], fila['matricula'], str(fila['nacimiento']))

    def _cursos(self, gestion, filas):
        cursos = {
            (c.grado, c.paralelo): c
            for c in Curso.objects.filter(gestion=gestion, nivel=Curso.Nivel.SECUNDARIA)
        }
        faltan = sorted({f['curso'] for f in filas.values()} - set(cursos))
        if faltan:
            raise CommandError(
                f'Estos cursos no existen en la gestión {gestion}: '
                + ', '.join(f'{g}° {p}' for g, p in faltan)
            )
        return cursos

    # ───────────────────────────── guardado ─────────────────────────────

    def _revisar_choques(self, vigentes):
        """Lo mismo que `EstudianteForm.clean_rude`: un RUDE no puede ser de otro."""
        rudes = [f['rude'] for f in vigentes]
        ajenos = list(
            Usuario.objects.filter(username__in=rudes).exclude(rol=Usuario.Rol.ESTUDIANTE)
            .values_list('username', flat=True)
        )
        tutores = list(
            Usuario.objects.filter(rol=Usuario.Rol.PADRE, ci__in=rudes).values_list('ci', flat=True)
        )
        if ajenos or tutores:
            raise CommandError(
                f'RUDE que ya usa otra cuenta: {ajenos} · RUDE igual al CI de un tutor: {tutores}'
            )

    def _guardar(self, vigentes, cursos, gestion):
        usuarios = {u.username.upper(): u for u in Usuario.objects.filter(rol=Usuario.Rol.ESTUDIANTE)}
        creados = actualizados = 0

        for fila in sorted(vigentes, key=lambda f: (f['curso'], f['nombre'])):
            nombres, apellidos = separar_nombre(fila['nombre'])
            usuario = usuarios.get(fila['rude'])
            nuevo = usuario is None
            if nuevo:
                usuario = Usuario(username=fila['rude'], rol=Usuario.Rol.ESTUDIANTE)

            cambio_carnet = nuevo or usuario.ci != fila['carnet']
            usuario.first_name, usuario.last_name = nombres, apellidos
            usuario.ci = fila['carnet']
            if cambio_carnet:
                # Cifrar es lento a propósito: solo cuando la contraseña cambia.
                usuario.set_password(fila['carnet'])
            usuario.save()

            nacimiento = fila['nacimiento']
            if isinstance(nacimiento, datetime.datetime):
                nacimiento = nacimiento.date()
            estudiante, _ = Estudiante.objects.update_or_create(
                usuario=usuario,
                defaults={
                    'rude': fila['rude'],
                    'fecha_nacimiento': nacimiento if isinstance(nacimiento, datetime.date) else None,
                    'genero': texto(fila['genero']).upper()[:1] if texto(fila['genero']).upper()[:1] in ('M', 'F') else '',
                    'nacionalidad': NACIONALIDADES.get(texto(fila['pais']).upper(), texto(fila['pais']).capitalize()),
                    'departamento': DEPARTAMENTOS.get(texto(fila['departamento']).upper(), ''),
                },
            )

            # Como `EstudianteForm._sincronizar_matricula`: una matrícula por año.
            curso = cursos[fila['curso']]
            Matricula.objects.filter(
                estudiante=estudiante, curso__gestion=gestion,
            ).exclude(curso=curso).update(activa=False)
            Matricula.objects.update_or_create(
                estudiante=estudiante, curso=curso, defaults={'activa': True},
            )

            if nuevo:
                creados += 1
            else:
                actualizados += 1
        return creados, actualizados

    def _desactivar_retirados(self, retirados, gestion):
        return Matricula.objects.filter(
            estudiante__rude__in=[f['rude'] for f in retirados],
            curso__gestion=gestion, activa=True,
        ).update(activa=False)

    # ───────────────────────────── resumen ─────────────────────────────

    def _resumen(self, vigentes, retirados, creados, actualizados, desactivados, cursos):
        por_curso = collections.Counter(f['curso'] for f in vigentes)
        self.stdout.write('\nCargados por curso:')
        for grado in sorted({g for g, _ in por_curso}):
            self.stdout.write('  ' + '   '.join(
                f'{grado}° {p}: {por_curso[(grado, p)]:>2}'
                for p in sorted(p for g, p in por_curso if g == grado)
            ))

        no_incorporados = sum(1 for f in vigentes if f['matricula'] == 'NO INCORPORADO')
        self.stdout.write(f'\nNO INCORPORADO cargados igual: {no_incorporados}')

        self.stdout.write(f'\nRetirados que no se cargaron ({len(retirados)}):')
        for f in sorted(retirados, key=lambda f: (f['curso'], f['nombre'])):
            self.stdout.write(f"  {f['curso'][0]}° {f['curso'][1]}  {f['rude']:<17} {f['nombre']}")
        if desactivados:
            self.stdout.write(f'  ({desactivados} ya estaban cargados: se desactivó su matrícula)')

        self.stdout.write(self.style.SUCCESS(
            f'\nListo: {creados} estudiantes nuevos, {actualizados} actualizados, '
            f'{len(cursos)} cursos en la gestión.'
        ))
